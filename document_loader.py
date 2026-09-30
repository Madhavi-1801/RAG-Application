"""
document_loader.py
===================
Handles loading documents of various file formats (.txt, .pdf, .docx) 
and extracting raw text content.

Extensibility:
--------------
New file handlers can be added to the self._handlers dictionary mapping file 
extensions (e.g. '.md', '.csv') to loader functions.
"""

import os
from typing import Callable, Dict

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import docx
except ImportError:
    docx = None


class DocumentLoader:
    """Loads text documents from local file paths based on file extension."""

    def __init__(self):
        # Register handlers for supported file formats.
        # To add a new file format (e.g., .md), define a helper method and register it here.
        self._handlers: Dict[str, Callable[[str], str]] = {
            ".txt": self._load_txt,
            ".pdf": self._load_pdf,
            ".docx": self._load_docx,
        }

    def load(self, file_path: str) -> str:
        """
        Loads the file at file_path and extracts its plain text content.
        
        Args:
            file_path (str): Path to the input file.
            
        Returns:
            str: Extracted text content.
            
        Raises:
            FileNotFoundError: If the file path does not exist.
            ValueError: If the file format is not supported or extraction fails.
        """
        # Clean up path quotes if user typed wrapped quotes (e.g. "C:\path\to\file.pdf")
        clean_path = file_path.strip("'\"")

        if not os.path.exists(clean_path):
            raise FileNotFoundError(f"File not found at path: '{clean_path}'")

        _, ext = os.path.splitext(clean_path)
        ext = ext.lower()

        if ext not in self._handlers:
            supported = ", ".join(self._handlers.keys())
            raise ValueError(
                f"Unsupported file format '{ext}'. Supported formats: {supported}"
            )

        handler = self._handlers[ext]
        text = handler(clean_path)

        if not text or not text.strip():
            raise ValueError(f"Extracted document text is empty: '{clean_path}'")

        return text.strip()

    def _load_txt(self, file_path: str) -> str:
        """Reads plain text files using UTF-8 encoding with fallback options."""
        encodings = ["utf-8", "latin-1", "cp1252"]
        for encoding in encodings:
            try:
                with open(file_path, "r", encoding=encoding) as f:
                    return f.read()
            except UnicodeDecodeError:
                continue
        raise ValueError(f"Could not decode text file with standard encodings: {file_path}")

    def _load_pdf(self, file_path: str) -> str:
        """Extracts text page-by-page from a PDF file using pypdf."""
        if PdfReader is None:
            raise ImportError(
                "pypdf is required to load PDF files. Install it using: pip install pypdf"
            )

        reader = PdfReader(file_path)
        pages_text = []

        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                pages_text.append(page_text)

        if not pages_text:
            raise ValueError(f"No readable text found in PDF: {file_path}")

        return "\n\n".join(pages_text)

    def _load_docx(self, file_path: str) -> str:
        """Extracts text from paragraphs and tables in a DOCX file using python-docx."""
        if docx is None:
            raise ImportError(
                "python-docx is required to load DOCX files. Install it using: pip install python-docx"
            )

        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also extract text from tables if present
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        return "\n\n".join(paragraphs)


if __name__ == "__main__":
    # Simple self-test
    loader = DocumentLoader()
    print("DocumentLoader initialized successfully. Supported extensions:", list(loader._handlers.keys()))
